"""DX（行政デジタル化）の集計：AI・RPA・生成AIの導入状況と、重点計画対象手続のオンライン化状況。

入力（download.sh で取得）
- soumu_ai.xlsx / soumu_rpa.xlsx / soumu_genai.xlsx：総務省「地方自治体におけるAI・RPAの実証実験・導入状況等調査」
  （令和7年度）の導入団体一覧。列：3 都道府県、4 団体名、5 導入済（○）、6 実証実験中（○）。
  一覧に掲載がない団体は「未導入」として扱う（未回答の団体を含む可能性がある）。
- soumu_dx_kobetsu.zip：総務省「自治体DX・情報化推進概要」個別資料（令和7年4月1日現在）の
  「（６）」ファイル・「市町村」シート。手続ごとの状況列（6, 11, 16, ... , 246）が「済」「未」。

出力（app.html の変数）：DX_ADMIN = {id: [AI, RPA, 生成AI（2=導入済・1=実証中・0=未導入）, オンライン化済の数, 対象の数]}
"""
import io
import os
import zipfile

import openpyxl

from common import DATA, NAME_TO_ID, js_block, save

LEVEL = {'導入済': 2, '実証中': 1, '未導入': 0}


def adoption(fname):
    ws = openpyxl.load_workbook(os.path.join(DATA, fname), read_only=True, data_only=True).active
    st = {}
    for r in ws.iter_rows(values_only=True):
        if len(r) > 6 and r[3] == '大阪府' and r[4] in NAME_TO_ID:
            st[NAME_TO_ID[r[4]]] = '導入済' if r[5] == '○' else ('実証中' if r[6] == '○' else '未導入')
    return st


def online():
    z = zipfile.ZipFile(os.path.join(DATA, 'soumu_dx_kobetsu.zip'))
    data = None
    for i in z.infolist():
        n = i.filename
        try:
            n = n.encode('cp437').decode('cp932')
        except UnicodeError:
            pass
        if '（６）' in n:
            data = z.read(i)
    ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)['市町村']
    cols = range(6, 251, 5)
    out = {}
    for r in ws.iter_rows(min_row=8, values_only=True):
        if not (r[1] and str(r[1]).startswith('27')):
            continue
        name = str(r[3]).replace('大阪府', '').strip()
        done = sum(1 for c in cols if r[c] == '済')
        applicable = sum(1 for c in cols if r[c] in ('済', '未'))
        out[NAME_TO_ID[name]] = (done, applicable)
    return out


def main():
    ai, rpa, genai = adoption('soumu_ai.xlsx'), adoption('soumu_rpa.xlsx'), adoption('soumu_genai.xlsx')
    onl = online()
    res = {}
    for i in NAME_TO_ID.values():
        res[i] = [LEVEL[ai.get(i, '未導入')], LEVEL[rpa.get(i, '未導入')], LEVEL[genai.get(i, '未導入')],
                  onl[i][0], onl[i][1]]
    save('dx_admin', res)
    print(js_block('DX_ADMIN', res, per_line=4))


if __name__ == '__main__':
    main()
